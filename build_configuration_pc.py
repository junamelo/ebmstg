from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "outputs" / "configuration_pc"
OUTPUT = OUTPUT_DIR / "Configuration_PC_Dell_Precision_7680_style_tableau.xlsx"


def build_workbook():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Caractéristiques PC"
    sheet.sheet_view.showGridLines = False

    # Titre du tableau, comme dans le modèle fourni.
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Tableau : Caractéristiques matérielles utilisées"
    sheet["A1"].font = Font(name="Arial", size=11, italic=True, color="007C92")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    sheet["A3"] = "Composant"
    sheet["B3"] = "Caractéristique"
    rows = [
        ["Marque et modèle", "Dell Precision 7680"],
        ["Processeur et fréquence", "Intel, environ 2,10 GHz"],
        ["Mémoire vive (RAM)", "Environ 32 Go"],
        ["Stockage", "Non indiqué dans les informations fournies"],
        ["Type du système", "Système 64 bits, processeur x64"],
        ["Système d'exploitation", "Windows 11 Professionnel"],
    ]
    for row_number, row in enumerate(rows, start=4):
        sheet.cell(row_number, 1, row[0])
        sheet.cell(row_number, 2, row[1])

    dark_teal = "15566B"
    light_blue = "DCEFF5"
    gray_blue = "D9D9D9"
    border_blue = "37A9D0"
    thin_blue = Side(style="thin", color=border_blue)
    border = Border(left=thin_blue, right=thin_blue, top=thin_blue, bottom=thin_blue)

    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor=dark_teal)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
    sheet.row_dimensions[3].height = 24

    for row_number in range(4, 10):
        label_cell = sheet.cell(row_number, 1)
        value_cell = sheet.cell(row_number, 2)
        label_cell.fill = PatternFill("solid", fgColor=light_blue if row_number % 2 == 0 else gray_blue)
        value_cell.fill = PatternFill("solid", fgColor=gray_blue if row_number % 2 == 0 else "FFFFFF")
        label_cell.font = Font(name="Arial", size=10, bold=True, color="000000")
        value_cell.font = Font(name="Arial", size=10, color="000000")
        for cell in (label_cell, value_cell):
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row_number].height = 30

    sheet.column_dimensions["A"].width = 31
    sheet.column_dimensions["B"].width = 72
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = "A3:B9"

    note_row = 11
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)
    sheet.cell(note_row, 1, "Source : informations systeminfo fournies pour le PC Dell Precision 7680")
    sheet.cell(note_row, 1).font = Font(name="Arial", size=9, italic=True, color="666666")
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True)

    workbook.properties.title = "Caractéristiques matérielles utilisées"
    workbook.properties.subject = "Configuration simplifiée du PC"
    workbook.properties.creator = "BANLEPO Benoît"
    workbook.save(OUTPUT)

    # Vérification de réouverture et des valeurs principales.
    check = load_workbook(OUTPUT, data_only=False)
    check_sheet = check["Caractéristiques PC"]
    assert check_sheet["A3"].value == "Composant"
    assert check_sheet["B3"].value == "Caractéristique"
    assert check_sheet["A4"].value == "Marque et modèle"
    assert check_sheet["B4"].value == "Dell Precision 7680"
    assert check_sheet.max_row == 11
    check.close()
    print(OUTPUT)


if __name__ == "__main__":
    build_workbook()

