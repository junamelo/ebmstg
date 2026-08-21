from copy import copy
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\Benoit\Downloads\ma\UCs.xlsx")
OUTPUT = Path("outputs") / "UCs_mis_a_jour.xlsx"


def copy_cell_style(source_cell, target_cell):
    """Copy only presentation attributes, never the source value."""
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)


def clear_merges_except_header(ws):
    for merged in list(ws.merged_cells.ranges):
        if str(merged) not in {"A1:B1", "A2:B2"}:
            ws.unmerge_cells(str(merged))


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Fichier source introuvable : {SOURCE}")

    wb = openpyxl.load_workbook(SOURCE)
    ws = wb["Feuil1"]

    # Save a small palette from the existing table so colors, fonts, borders,
    # alignment and wrapping stay identical to the original workbook.
    palette = {
        "a": copy(ws["A5"]),
        "b": copy(ws["B5"]),
        "c": copy(ws["C5"]),
        "a_mid": copy(ws["A6"]),
        "b_mid": copy(ws["B6"]),
        "c_mid": copy(ws["C6"]),
        "a_bottom": copy(ws["A9"]),
        "b_bottom": copy(ws["B9"]),
        "c_bottom": copy(ws["C9"]),
    }

    clear_merges_except_header(ws)
    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row - 2)

    # The table is deliberately kept in the same three-column presentation as
    # the reference: family of use cases, detailed use case and actors.
    groups = [
        (
            "Authentification et profil",
            [
                ("Se connecter à la plateforme", "Super Administrateur, Chef de facturation, Agent de facturation, Commercial, Payeur, Employé"),
                ("Réinitialiser le mot de passe", "Super Administrateur, Chef de facturation, Agent de facturation, Commercial, Payeur, Employé"),
                ("Gérer son profil et activer/désactiver le 2FA", "Super Administrateur, Chef de facturation, Agent de facturation, Commercial, Payeur, Employé"),
            ],
        ),
        (
            "Gestion des comptes et des rôles",
            [
                ("Créer un compte utilisateur", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Modifier les informations d'un compte", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Consulter la liste et le détail des comptes", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Activer ou désactiver un compte", "Super Administrateur"),
                ("Gérer les rôles et les droits d'accès", "Super Administrateur"),
            ],
        ),
        (
            "Gestion des acteurs métier",
            [
                ("Créer et consulter les commerciaux", "Super Administrateur, Chef de facturation"),
                ("Créer et gérer les payeurs", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Créer un employé et l'associer à une entreprise et à une ligne", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Associer plusieurs lignes existantes à un payeur", "Super Administrateur, Chef de facturation, Agent de facturation"),
            ],
        ),
        (
            "Gestion des demandes et des contrats",
            [
                ("Soumettre une demande de contrat", "Commercial"),
                ("Consulter le détail et suivre l'état d'une demande", "Commercial, Chef de facturation, Agent de facturation"),
                ("Valider ou rejeter une demande avec un motif", "Chef de facturation, Agent de facturation"),
                ("Créer et modifier les informations d'un contrat validé", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Gérer les lignes, options et services prévus au contrat", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Résilier un contrat, renseigner le motif et la date", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Consulter l'historique des actions et exporter le contrat en PDF", "Super Administrateur, Chef de facturation, Agent de facturation"),
            ],
        ),
        (
            "Gestion des services, forfaits et tarifs",
            [
                ("Créer ou modifier un type de service", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Créer, modifier ou désactiver un service ou une option", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Créer et gérer les forfaits, leurs options et leurs prix", "Super Administrateur, Chef de facturation, Agent de facturation"),
                ("Configurer les règles tarifaires de la simulation", "Super Administrateur, Chef de facturation, Agent de facturation"),
            ],
        ),
        (
            "Importation et publication des factures PDF",
            [
                ("Importer un bloc PDF global ou un bloc de factures sommaires", "Chef de facturation, Agent de facturation"),
                ("Découper automatiquement le bloc et associer chaque facture au bon compte", "Chef de facturation, Agent de facturation"),
                ("Consulter les factures à publier et lancer la publication", "Chef de facturation, Agent de facturation"),
                ("Notifier les utilisateurs de la disponibilité des factures par e-mail", "Chef de facturation, Agent de facturation"),
                ("Consulter l'historique et le rapport des publications", "Chef de facturation, Agent de facturation"),
            ],
        ),
        (
            "Consultation des factures",
            [
                ("Consulter et filtrer ses factures par période", "Payeur, Employé"),
                ("Consulter la facture globale et les factures sommaires de la flotte", "Payeur"),
                ("Consulter la facture sommaire de sa ligne", "Employé"),
                ("Ouvrir dans un nouvel onglet ou télécharger le PDF publié", "Payeur, Employé"),
            ],
        ),
        (
            "Simulation de facturation",
            [
                ("Simuler le montant de la facturation d'un contrat", "Payeur"),
                ("Simuler le montant de la facturation de sa ligne", "Employé"),
                ("Consulter l'historique des simulations", "Payeur, Employé"),
            ],
        ),
    ]

    row = 3
    for category, items in groups:
        start = row
        for index, (use_case, actors) in enumerate(items):
            # Reuse the original style pattern for the first, middle and last
            # row of each category. This leaves the existing visual language
            # untouched while allowing the updated table to grow naturally.
            if index == 0:
                a_style, b_style, c_style = palette["a"], palette["b"], palette["c"]
            elif index == len(items) - 1:
                a_style, b_style, c_style = palette["a_bottom"], palette["b_bottom"], palette["c_bottom"]
            else:
                a_style, b_style, c_style = palette["a_mid"], palette["b_mid"], palette["c_mid"]

            copy_cell_style(a_style, ws.cell(row, 1))
            copy_cell_style(b_style, ws.cell(row, 2))
            copy_cell_style(c_style, ws.cell(row, 3))
            ws.cell(row, 1).value = category if index == 0 else None
            ws.cell(row, 2).value = use_case
            ws.cell(row, 3).value = actors

            # Give long descriptions enough vertical space without changing
            # the column widths, fills, fonts or borders of the source.
            longest = max(len(use_case), len(actors))
            ws.row_dimensions[row].height = 60.6 if longest > 75 else (45.6 if longest > 45 else 30.6)
            row += 1

        ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Créé : {OUTPUT.resolve()}")
    print(f"Lignes de cas d'utilisation : {row - 3}")


if __name__ == "__main__":
    main()
