from pathlib import Path
from copy import copy

import openpyxl
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("outputs") / "UCs_mis_a_jour.xlsx"
OUTPUT = Path("outputs") / "UCs_mis_a_jour_Arial12.docx"


def set_run_font(run, name="Arial", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_cell_text(cell, value, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("" if value is None else str(value))
    set_run_font(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="808080", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Feuil1"]

    rows = []
    for row_index in range(2, ws.max_row + 1):
        rows.append([ws.cell(row_index, col).value for col in range(1, 4)])

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    run = caption.add_run("Tableau des cas d’utilisation – Moov e-Factures")
    set_run_font(run, bold=True)

    table = document.add_table(rows=len(rows), cols=3)
    table.autofit = False
    table.allow_autofit = False
    widths_cm = [3.4, 8.3, 6.3]
    set_table_widths(table, widths_cm)

    # Header: keep the same logical grouping as the Excel source.
    header = table.rows[0]
    header.cells[0].merge(header.cells[1])
    set_cell_text(header.cells[0], "CAS D’UTILISATION", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(header.cells[2], "ACTEURS", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in header.cells:
        set_cell_margins(cell, top=100, bottom=100)
        set_cell_borders(cell, color="404040", size="8")
    set_repeat_table_header(header)
    set_no_split(header)

    # Body values and formatting.
    for i, values in enumerate(rows[1:], start=1):
        row = table.rows[i]
        set_cell_text(row.cells[0], values[0], bold=values[0] is not None)
        set_cell_text(row.cells[1], values[1])
        set_cell_text(row.cells[2], values[2])
        for cell in row.cells:
            set_cell_margins(cell)
            set_cell_borders(cell)
        set_no_split(row)

    # Recreate the category merges from the updated Excel file.
    for merged in ws.merged_cells.ranges:
        if merged.min_col == 1 and merged.max_col == 1 and merged.min_row >= 3:
            # Excel row 3 maps to Word table row index 1.
            start = merged.min_row - 2
            end = merged.max_row - 2
            if end > start:
                merged_cell = table.cell(start, 0).merge(table.cell(end, 0))
                set_cell_text(merged_cell, ws.cell(merged.min_row, 1).value, bold=True)
                set_cell_margins(merged_cell)
                set_cell_borders(merged_cell)

    # Ensure the entire document (including merged cells) is Arial 12.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, size=12, bold=run.bold is True)
    for table_obj in document.tables:
        for row in table_obj.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=12, bold=run.bold is True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Créé : {OUTPUT.resolve()}")
    print(f"Table : {len(rows) - 1} cas d’utilisation, police Arial 12")


if __name__ == "__main__":
    main()
