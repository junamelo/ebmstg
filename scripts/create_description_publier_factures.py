from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Benoit\Downloads\ma\lesDescri.xlsx")
OUTPUT = ROOT / "docs" / "description_textuelle_publier_factures.xlsx"

wb = load_workbook(TEMPLATE)
ws = wb.active
ws.title = "Publier les factures"

# Keep the first example's visual language, but make this workbook contain one
# independent two-column description table.
for merged in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(merged))
ws.delete_cols(3, ws.max_column - 2)

# Remove the previous example's text while preserving its fonts, fills and borders.
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
    for cell in row:
        cell.value = None
        cell.alignment = copy(cell.alignment)
        cell.alignment = Alignment(
            horizontal=cell.alignment.horizontal or "left",
            vertical="center",
            wrap_text=True,
            text_rotation=cell.alignment.text_rotation,
            shrink_to_fit=cell.alignment.shrink_to_fit,
            indent=cell.alignment.indent,
        )

# Header and section layout. The merged ranges reproduce the reference logic:
# one label in column A, with its multi-line content continuing in column B.
for merge in [
    "A5:A6",
    "A8:A9",
    "A10:A11",
    "A12:A13",
    "A14:A27",
    "A28:A43",
    "A44:A46",
]:
    ws.merge_cells(merge)

values = {
    "A1": "Élément",
    "B1": "Description",
    "A2": "Titre",
    "B2": "Publier les factures",
    "A4": "Acteur principal",
    "B4": "Agent de facturation",
    "A5": "Résumé",
    "B5": "Ce cas d'utilisation décrit la mise à disposition des factures clients postpayés après le traitement des blocs PDF. L'agent vérifie les factures prêtes à publier, lance la publication et contrôle le résultat de l'opération.",
    "A7": "Version",
    "B7": "1.0",
    "A8": "Date de création",
    "B8": datetime(2026, 8, 12),
    "A10": "Date de modification",
    "B10": datetime(2026, 8, 12),
    "A12": "Précondition",
    "B12": "• L'agent de facturation est authentifié et possède le droit de publier les factures.\n• Les factures ont été extraites du bloc PDF et sont associées à un contrat ou à une ligne.\n• Les fichiers PDF nécessaires sont disponibles dans le stockage média.",
    "A14": "Scénario nominal",
    "A28": "Scénario alternatif",
    "A44": "Scénarios d'exceptions",
    "A47": "Postcondition",
    "B47": "Les factures publiées sont enregistrées avec leur statut, leur période, leur numéro, leur montant et leur fichier PDF. Elles deviennent consultables par le payeur ou l'employé concerné. L'historique de publication est mis à jour.",
}

nominal = [
    "1. L'agent de facturation ouvre la page des factures à publier.",
    "2. Le système affiche les factures prêtes à être publiées avec leur numéro, leur type, leur période et leur fichier PDF.",
    "3. L'agent vérifie les informations affichées et peut filtrer les factures globales ou sommaires.",
    "4. L'agent sélectionne les factures à publier.",
    "5. L'agent clique sur « Publier les factures ».",
    "6. Le système vérifie que chaque facture possède un fichier PDF et qu'elle n'a pas déjà été publiée.",
    "7. Le système place le traitement de publication en file d'attente afin de traiter le bloc sans bloquer l'interface.",
    "8. Le traitement associe chaque fichier PDF à la facture correspondante.",
    "9. Le système met à jour le statut des factures publiées.",
    "10. Le système enregistre la date, la période et l'agent ayant lancé la publication.",
    "11. Le système met à jour l'historique des publications.",
    "12. Si l'option e-mail est activée, le système prépare une notification de disponibilité de la facture.",
    "13. Le système affiche le nombre de PDF créés, de factures mises à jour et les éventuelles erreurs.",
    "14. L'agent peut consulter le détail du traitement et revenir à l'historique.",
]
for index, text in enumerate(nominal, start=14):
    values[f"B{index}"] = text

alternatives = {
    "B28": "A1. Facture déjà traitée",
    "B29": "• Le système détecte que la facture est déjà publiée ou déjà traitée.",
    "B30": "• La facture n'est pas retraitée et apparaît dans le bilan avec un avertissement orange.",
    "B31": "A2. Fichier PDF manquant",
    "B32": "• Le système ne publie pas la facture concernée et signale l'absence du fichier PDF en rouge.",
    "B33": "• Les autres factures valides continuent leur traitement.",
    "B34": "A3. Correspondance facture-utilisateur incomplète",
    "B35": "• Le système conserve la facture en attente et indique qu'aucun contrat ou aucune ligne n'a pu être identifié.",
    "B36": "A4. Publication partielle",
    "B37": "• Lorsque certaines factures réussissent et d'autres échouent, le système conserve les résultats séparément.",
    "B38": "• Le bilan indique précisément les factures publiées, ignorées et en erreur.",
    "B39": "A5. Notification e-mail sélectionnée",
    "B40": "• Le système prépare une notification pour les destinataires disposant d'une adresse e-mail valide.",
    "B41": "• Une facture reste publiée même si l'envoi de la notification échoue.",
    "B42": "A6. Aucun élément sélectionné",
    "B43": "• Le système demande à l'agent de sélectionner au moins une facture avant de lancer la publication.",
}
values.update(alternatives)

exceptions = {
    "B44": "E1. Problème de connexion : si l'interface ne peut plus joindre le serveur, le système affiche un message d'erreur et aucun nouveau traitement n'est lancé.",
    "B45": "E2. Service de traitement indisponible : si le worker Celery ou le broker Garnet est arrêté, le traitement reste en attente et l'agent est informé.",
    "B46": "E3. Erreur de stockage ou de base de données : le système conserve le détail de l'erreur, marque la facture en échec et évite de créer une publication incohérente.",
}
values.update(exceptions)

for coordinate, value in values.items():
    ws[coordinate] = value

# Keep the reference's dimensions and make all descriptions readable.
ws.column_dimensions["A"].width = 28.109375
ws.column_dimensions["B"].width = 70
for row in range(1, 48):
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 30)
for row in [4, 12, 47]:
    ws.row_dimensions[row].height = 90
for row in range(14, 28):
    ws.row_dimensions[row].height = 42
for row in range(28, 44):
    ws.row_dimensions[row].height = 42
for row in range(44, 47):
    ws.row_dimensions[row].height = 75

for row in ws.iter_rows(min_row=1, max_row=47, min_col=1, max_col=2):
    for cell in row:
        cell.font = copy(cell.font)
        cell.font = cell.font.copy(name="Arial", sz=12)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal or "left", vertical="center", wrap_text=True)

ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
