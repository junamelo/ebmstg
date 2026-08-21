"""Create the Word version of the textual use-case description.

The layout follows the two-column structure of the reference workbook.  The
first column contains the section label and the second column contains the
description or the individual scenario steps.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "description_textuelle_publier_factures.docx"

ARIAL = "Arial"
TABLE_WIDTH_DXA = 14900
LABEL_WIDTH_DXA = 4000
DESCRIPTION_WIDTH_DXA = TABLE_WIDTH_DXA - LABEL_WIDTH_DXA


def set_run_font(run, *, size=12, bold=False, color="000000", italic=False):
    run.font.name = ARIAL
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), ARIAL)
    color_el = rpr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        rpr.append(color_el)
    color_el.set(qn("w:val"), color)


def clear_cell(cell):
    cell.text = ""
    # python-docx leaves one paragraph in a cell after assigning an empty text.
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    return paragraph


def set_cell_text(cell, value, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = clear_cell(cell)
    paragraph.alignment = align
    lines = str(value or "").split("\n")
    first = True
    for line in lines:
        target = paragraph if first else cell.add_paragraph()
        first = False
        target.alignment = align
        target.paragraph_format.space_before = Pt(0)
        target.paragraph_format.space_after = Pt(0)
        target.paragraph_format.line_spacing = 1.05
        run = target.add_run(line)
        set_run_font(run, bold=bold)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="808080", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table):
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "808080")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in (LABEL_WIDTH_DXA, DESCRIPTION_WIDTH_DXA):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_no_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, LABEL_WIDTH_DXA if index == 0 else DESCRIPTION_WIDTH_DXA)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)


def build_content_rows():
    rows = [["", ""] for _ in range(47)]

    def put(row_number, label=None, description=None):
        row = rows[row_number - 1]
        if label is not None:
            row[0] = label
        if description is not None:
            row[1] = description

    put(2, "Titre", "Publier les factures")
    put(4, "Acteur principal", "Agent de facturation")
    put(
        5,
        "Résumé",
        "Ce cas d'utilisation décrit la mise à disposition des factures clients postpayés après le traitement des blocs PDF. L'agent vérifie les factures prêtes à publier, lance la publication et contrôle le résultat de l'opération.",
    )
    put(7, "Version", "1.0")
    put(8, "Date de création", "12/08/2026")
    put(10, "Date de modification", "12/08/2026")
    put(
        12,
        "Précondition",
        "L'agent de facturation est authentifié et possède le droit de publier les factures ; les factures ont été extraites du bloc PDF et sont associées à un contrat ou à une ligne ; les fichiers PDF nécessaires sont disponibles dans le stockage média.",
    )
    put(14, "Scénario nominal")

    nominal = [
        "1. L'agent de facturation ouvre la page des factures à publier.",
        "2. Le système affiche les factures prêtes à être publiées avec leur numéro, leur type, leur période et leur fichier PDF.",
        "3. L'agent vérifie les informations affichées et peut filtrer les factures globales ou sommaires.",
        "4. L'agent sélectionne les factures à publier.",
        "5. L'agent clique sur « Publier les factures ».",
        "6. Le système vérifie que chaque facture possède un fichier PDF et qu'elle n'a pas déjà été publiée.",
        "7. Le système place le traitement de publication en file d'attente afin de traiter le bloc sans bloquer l'interface.",
        "8. Le traitement associe chaque fichier PDF à la facture correspondante.",
        "9. Le système met à jour le statut des factures publiées.",
        "10. Le système enregistre la date, la période et l'agent ayant lancé la publication.",
        "11. Le système met à jour l'historique des publications.",
        "12. Si l'option e-mail est activée, le système prépare une notification de disponibilité de la facture.",
        "13. Le système affiche le nombre de PDF créés, de factures mises à jour et les éventuelles erreurs.",
        "14. L'agent peut consulter le détail du traitement et revenir à l'historique.",
    ]
    for row_number, text in enumerate(nominal, start=14):
        put(row_number, description=text)

    put(28, "Scénario alternatif", "A1. Facture déjà traitée")
    put(29, description="Le système détecte que la facture est déjà publiée ou déjà traitée.")
    put(30, description="La facture n'est pas retraitée et apparaît dans le bilan avec un avertissement orange.")
    put(31, description="A2. Fichier PDF manquant")
    put(32, description="Le système ne publie pas la facture concernée et signale l'absence du fichier PDF en rouge.")
    put(33, description="Les autres factures valides continuent leur traitement.")
    put(34, description="A3. Correspondance facture-utilisateur incomplète")
    put(35, description="Le système conserve la facture en attente et indique qu'aucun contrat ou aucune ligne n'a pu être identifié.")
    put(36, description="A4. Publication partielle")
    put(37, description="Lorsque certaines factures réussissent et d'autres échouent, le système conserve les résultats séparément.")
    put(38, description="Le bilan indique précisément les factures publiées, ignorées et en erreur.")
    put(39, description="A5. Notification e-mail sélectionnée")
    put(40, description="Le système prépare une notification pour les destinataires disposant d'une adresse e-mail valide.")
    put(41, description="Une facture reste publiée même si l'envoi de la notification échoue.")
    put(42, description="A6. Aucun élément sélectionné")
    put(43, description="Le système demande à l'agent de sélectionner au moins une facture avant de lancer la publication.")

    put(44, "Scénarios d'exceptions", "E1. Problème de connexion : si l'interface ne peut plus joindre le serveur, le système affiche un message d'erreur et aucun nouveau traitement n'est lancé.")
    put(45, description="E2. Service de traitement indisponible : si le worker Celery ou le broker Garnet est arrêté, le traitement reste en attente et l'agent est informé.")
    put(46, description="E3. Erreur de stockage ou de base de données : le système conserve le détail de l'erreur, marque la facture en échec et évite de créer une publication incohérente.")
    put(47, "Postcondition", "Les factures publiées sont enregistrées avec leur statut, leur période, leur numéro, leur montant et leur fichier PDF. Elles deviennent consultables par le payeur ou l'employé concerné. L'historique de publication est mis à jour.")
    return rows


def main():
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    normal = document.styles["Normal"]
    normal.font.name = ARIAL
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), ARIAL)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), ARIAL)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), ARIAL)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("Description textuelle du cas d'utilisation")
    set_run_font(title_run, size=14, bold=True, color="1F4E78")

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption_run = caption.add_run("Cas d'utilisation : Publier les factures")
    set_run_font(caption_run, size=12, bold=True, color="404040")

    # One coherent table, with the same two-column arrangement as the reference.
    rows = build_content_rows()
    table = document.add_table(rows=len(rows), cols=2)
    table.cell(0, 0).text = "Élément"
    table.cell(0, 1).text = "Description"

    # Excel row 1 is the Word header; Excel rows 2..47 map to Word rows 1..46.
    for excel_row in range(2, 48):
        word_row = excel_row - 1
        label, description = rows[excel_row - 1]
        set_cell_text(table.cell(word_row, 0), label, bold=bool(label))
        set_cell_text(table.cell(word_row, 1), description)

    # The header is the first Excel row. Content rows below it correspond to rows 2..47.
    set_cell_text(table.cell(0, 0), "Élément", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), "Description", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "1F4E78")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=12, bold=True, color="FFFFFF")

    for first, last in ((5, 6), (8, 9), (10, 11), (12, 13), (14, 27), (28, 43), (44, 46)):
        merged = table.cell(first - 1, 0).merge(table.cell(last - 1, 0))
        set_cell_text(merged, rows[first - 1][0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(merged, "D9EAF7")

    # Labels that are not part of a merged group use the same restrained fill.
    for row_number in (2, 4, 7, 47):
        set_cell_shading(table.cell(row_number - 1, 0), "D9EAF7")

    set_table_geometry(table)

    # Re-apply table text fonts after merging cells, which can create new runs.
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, size=12, bold=run.bold is True, color="FFFFFF" if row is table.rows[0] else "000000")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
