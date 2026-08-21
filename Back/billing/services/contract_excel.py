"""Génération du fichier Excel de synthèse d'un contrat."""

from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BLUE = '002A7A'
LIGHT_BLUE = 'EAF2FF'
LIGHT_GREY = 'F7F7F7'
WHITE = 'FFFFFF'
GRID = 'D1D5DB'


def _value(value):
    """Retourne une valeur lisible sans afficher None dans Excel."""
    return '-' if value in (None, '') else value


def _employee_name(line):
    if not line.employe:
        return '-'
    name = f'{line.employe.first_name} {line.employe.last_name}'.strip()
    return name or line.employe.username


def _logo_path():
    path = Path(settings.BASE_DIR).parent / 'Front' / 'src' / 'assets' / 'logo-moov.png'
    return path if path.is_file() else None


def _add_logo(sheet):
    """Ajoute le logo dans l'en-tête lorsque le fichier est présent."""
    path = _logo_path()
    if not path:
        return
    logo = ExcelImage(str(path))
    logo.width = 135
    logo.height = 59
    sheet.add_image(logo, 'A1')


def _set_title(sheet, title, subtitle):
    sheet.merge_cells('A1:B3')
    sheet.merge_cells('C1:F2')
    sheet.merge_cells('C3:F3')
    for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=6):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=BLUE)
    sheet['C1'] = title
    sheet['C1'].font = Font(name='Arial', size=16, bold=True, color=WHITE)
    sheet['C1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet['C3'] = subtitle
    sheet['C3'].font = Font(name='Arial', size=10, color='DBEAFE')
    sheet['C3'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 21
    _add_logo(sheet)


def _thin_border():
    side = Side(style='thin', color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_information_table(sheet, start_row, rows, label_col, value_col):
    label_fill = PatternFill('solid', fgColor=LIGHT_BLUE)
    for row_index, (label, value) in enumerate(rows, start=start_row):
        label_cell = sheet.cell(row_index, label_col, label)
        value_cell = sheet.cell(row_index, value_col, _value(value))
        label_cell.font = Font(name='Arial', size=10, bold=True, color=BLUE)
        label_cell.fill = label_fill
        for cell in (label_cell, value_cell):
            if cell is value_cell:
                cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = _thin_border()


def _service_status(line, field):
    return 'Oui' if getattr(line, field) else 'Non'


def generer_excel_contrat(company):
    """Retourne un BytesIO avec les informations et les lignes du contrat."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Contrat'
    lines_sheet = workbook.create_sheet('Lignes')

    _set_title(summary, 'MOOV AFRICA - CONTRAT CLIENT', f'Code contrat : {company.compte}')
    summary.sheet_view.showGridLines = False
    summary.column_dimensions['A'].width = 24
    summary.column_dimensions['B'].width = 42
    summary.column_dimensions['C'].width = 4
    summary.column_dimensions['D'].width = 24
    summary.column_dimensions['E'].width = 42
    summary.column_dimensions['F'].width = 4

    payeur = '-'
    if company.payeur:
        payeur = f'{company.payeur.first_name} {company.payeur.last_name}'.strip()
        if company.payeur.email:
            payeur = f'{payeur} - {company.payeur.email}'
    commercial = '-'
    if company.commercial:
        commercial = f'{company.commercial.prenom} {company.commercial.nom} ({company.commercial.matricule})'

    left_rows = [
        ('Raison sociale', company.raison_sociale),
        ('Catégorie', company.get_categorie_display()),
        ('Statut contrat', company.statut),
        ('Statut factures', company.get_statut_factures_display()),
        ('Date effet', company.date_effet),
        ('Date fin', company.date_fin),
        ('Mode de règlement', company.get_mode_reglement_display()),
        ('Exonéré de TVA', 'Oui' if company.est_exonere else 'Non'),
        ('Adresse 1', company.adresse),
        ('Adresse 2', company.adresse_ligne2),
    ]
    right_rows = [
        ('Payeur', payeur),
        ('Commercial', commercial),
        ('E-mail de facturation', company.email_facturation),
        ('Type de revenu', company.type_revenu),
        ('Observation', company.observation),
        ('Contrat résilié', 'Oui' if company.est_resilie else 'Non'),
        ('Date de résiliation', company.date_resiliation),
        ('Motif de résiliation', company.motif_resiliation),
        ('Observation résiliation', company.observation_resiliation),
        ('Date de génération', timezone.localtime().strftime('%d/%m/%Y %H:%M')),
    ]
    _style_information_table(summary, 5, left_rows, 1, 2)
    _style_information_table(summary, 5, right_rows, 4, 5)
    summary.freeze_panes = 'A5'
    summary.print_area = 'A1:E14'
    summary.page_setup.orientation = 'landscape'
    summary.page_setup.fitToWidth = 1

    _set_title(lines_sheet, 'MOOV AFRICA - LIGNES DU CONTRAT', f'Contrat : {company.compte}')
    lines_sheet.sheet_view.showGridLines = False
    headers = [
        'N°', 'MSISDN', 'Employé', 'E-mail employé', 'Statut', 'Cycle',
        'Forfait (FCFA)', 'Fact. détaillée', 'No Limit', 'BlackBerry',
        'Incognito', 'Roaming', 'Internet', 'International', 'Non revenu',
    ]
    header_row = 5
    for column, header in enumerate(headers, start=1):
        cell = lines_sheet.cell(header_row, column, header)
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.font = Font(name='Arial', size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    lines = company.lines.select_related('employe').order_by('msisdn')
    for index, line in enumerate(lines, start=1):
        row = header_row + index
        values = [
            index, line.msisdn, _employee_name(line),
            line.employe.email if line.employe else '-', line.statut, line.cycle,
            line.forfait, _service_status(line, 'facture_detaillee'),
            _value(line.option_nolimit), _value(line.option_blackberry),
            _service_status(line, 'est_incognito'), _service_status(line, 'est_roaming'),
            _service_status(line, 'est_internet'), _service_status(line, 'est_international'),
            _service_status(line, 'est_non_revenu'),
        ]
        for column, value in enumerate(values, start=1):
            cell = lines_sheet.cell(row, column, value)
            cell.font = Font(name='Arial', size=9)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = PatternFill('solid', fgColor=WHITE if index % 2 else LIGHT_GREY)
            cell.border = Border(bottom=Side(style='thin', color='E5E7EB'))
        lines_sheet.cell(row, 7).number_format = '#,##0 "FCFA"'

    last_row = max(header_row, header_row + company.lines.count())
    lines_sheet.auto_filter.ref = f'A{header_row}:O{last_row}'
    lines_sheet.freeze_panes = 'A6'
    widths = [7, 15, 25, 30, 14, 12, 18, 16, 16, 16, 12, 12, 12, 16, 16]
    for column, width in enumerate(widths, start=1):
        lines_sheet.column_dimensions[get_column_letter(column)].width = width
    lines_sheet.row_dimensions[header_row].height = 34
    lines_sheet.page_setup.orientation = 'landscape'
    lines_sheet.page_setup.fitToWidth = 1
    lines_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    lines_sheet.print_title_rows = f'1:{header_row}'

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
